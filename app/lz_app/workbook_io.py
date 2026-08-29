"""Workbook I/O for the app: values-aware xlsx export + import.

Export reuses lz_spec/gen_template.py's emitters (so the produced workbook is
a REAL template - banner rows, sentinels, type hints, dropdowns, derived
striping) with the IR's values written into the entry cells:
  - object tables: rows become the value-cell rows (the same mechanism the
    template uses for sample rows);
  - scalar tables: the Value column is prefilled (the _meta mechanism, applied
    to every sheet);
  - list tables: values land in the Value column.

import_workbook() reads such a workbook back to the identical spec dict, so
IR -> xlsx -> IR round-trips exactly for schema-clean specs (gate-tested).
Content the schema does not know (legacy extra tables/fields) is dropped on
export and reported.
"""

import copy
from pathlib import Path


def export_workbook(ir: dict, path: Path) -> list:
    """Write ir['sheets'] into a full template workbook. Returns notes about
    content the schema doesn't cover (dropped)."""
    from lz_spec import gen_template as gt
    from lz_spec import schema as wb_schema
    from openpyxl import Workbook

    sheets_data = ir.get("sheets", {})
    notes = []
    known = {s.name for s in wb_schema.SHEETS}
    for sname in sheets_data:
        if sname not in known and sname not in wb_schema.INFO_SHEETS:
            notes.append(f"sheet {sname!r} is not in the schema - not exported")

    data_sheets = []
    for sdef in wb_schema.SHEETS:
        if sdef.name in wb_schema.INFO_SHEETS:
            data_sheets.append(sdef)  # emit template content as-is
            continue
        sdata = sheets_data.get(sdef.name) or {}
        tables = []
        tnames = {t.name for t in sdef.tables}
        for tname in sdata:
            if tname not in tnames:
                notes.append(f"{sdef.name}.{tname}: unknown table - not exported")
        for tdef in sdef.tables:
            t = copy.copy(tdef)
            tdata = sdata.get(tdef.name)
            if tdef.kind == "scalar":
                values = tdata or {}
                for f in values:
                    if f not in {kv.name for kv in tdef.rows}:
                        notes.append(f"{sdef.name}.{tdef.name}.{f}: unknown field - not exported")
                t.rows = [gt.KV(kv.name, kv.type, values.get(kv.name), kv.sample, kv.description)
                          for kv in tdef.rows]
            elif tdef.kind == "list-single":
                t.sample_rows = list(tdata or [])
            else:
                t.sample_rows = [dict(r) for r in (tdata or [])]
            tables.append(t)
        s = copy.copy(sdef)
        s.tables = tables
        data_sheets.append(s)

    wb = Workbook()
    wb.remove(wb.active)
    gt._COL_POS.clear()
    orig_scalar = gt._emit_scalar
    # every scalar table gets its Value column prefilled (KV.default = IR value)
    gt._emit_scalar = lambda ws, r, t, prefill=False: orig_scalar(ws, r, t, prefill=True)
    try:
        for sheet in data_sheets:
            gt._emit_sheet(wb, sheet)
        gt._add_validations(wb)
    finally:
        gt._emit_scalar = orig_scalar
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return notes


def import_workbook(path: Path) -> dict:
    """xlsx -> IR (same shape lz_pipeline.model produces)."""
    from lz_pipeline import model
    ir, warnings = model.from_workbook(Path(path))
    ir["_import_warnings"] = warnings
    return ir


def normalize_ir(ir: dict) -> dict:
    """Canonical form: every schema column/field present (None when absent),
    blank strings as None. Round-trip identity holds on this form:
    normalize(ir) == import(export(ir)). The app normalizes on load so the
    editor always shows the full column set."""
    from lz_spec import schema as wb_schema
    out = copy.deepcopy(ir)
    sheets = out.setdefault("sheets", {})

    def _norm(v, typ=None):
        if isinstance(v, str) and not v.strip():
            return None
        # object-table bool cells may be real Excel booleans OR "TRUE"/"FALSE"
        # strings depending on how the cell was entered; canonical form is the
        # string (what the template writer emits, identical to _truthy()).
        if typ == "bool" and isinstance(v, bool):
            return "TRUE" if v else "FALSE"
        return v

    for sdef in wb_schema.SHEETS:
        if sdef.name in wb_schema.INFO_SHEETS:
            continue
        sdata = sheets.setdefault(sdef.name, {})
        for tdef in sdef.tables:
            if tdef.kind == "scalar":
                t = sdata.setdefault(tdef.name, {})
                for kv in tdef.rows:
                    t[kv.name] = _norm(t.get(kv.name))
            elif tdef.kind == "list-single":
                sdata.setdefault(tdef.name, [])
            else:
                rows = sdata.setdefault(tdef.name, [])
                cols = [("Enabled", "bool")] * (not tdef.mandatory and not any(
                    c[0] == "Enabled" for c in tdef.columns)) + \
                    [(c[0], c[1]) for c in tdef.columns]
                for row in rows:
                    for cname, ctype in cols:
                        row[cname] = _norm(row.get(cname), ctype)
                # all-blank rows can't survive a workbook round-trip (the
                # importer skips them), so canonical form drops them too
                sdata[tdef.name] = [r for r in rows
                                    if any(v is not None for v in r.values())]
    return out
