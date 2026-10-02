"""gen_ipam --states-dir: VPC notes and subnet purposes describe what is deployed.

Builds a two-VPC estate (a hub DMZ and one spoke) as 05-network tfvars plus
state pulls, and checks the remarks name the hosted resources, the workload
tags of VMs, the empty subnets, and VPC-level features - and that without
--states-dir the generator keeps its short role labels.
"""

import json
import subprocess
import sys
from pathlib import Path

import openpyxl

PIPELINE = Path(__file__).resolve().parents[2] / "pipeline"


def _estate(tmp: Path) -> tuple[Path, Path]:
    envs, states = tmp / "envs", tmp / "states"
    (envs / "05-network").mkdir(parents=True)
    states.mkdir()
    (envs / "05-network" / "terraform.tfvars.json").write_text(json.dumps({
        "spoke_private_supernet": "10.0.0.0/16",
        "hub_account": "lz-infra",
        "hub_vpcs": {"lz-hub-vpc-dmz-01": {"cidr": "10.0.0.0/22", "subnets": [
            {"name": "lz-hub-subnet-nat-01", "cidr": "10.0.0.0/25"},
            {"name": "lz-hub-subnet-elb-01", "cidr": "10.0.0.128/25"}]}},
        "spokes": {"lz-app-vpc-01": {"account": "lz-app", "vpc_cidr": "10.0.4.0/22", "subnets": [
            {"name": "lz-app-subnet-compute-01", "cidr": "10.0.4.0/25"},
            {"name": "lz-app-subnet-att-01", "cidr": "10.0.7.240/28"}]}},
    }), encoding="utf-8")

    def res(t, *attrs):
        return {"mode": "managed", "type": t, "instances": [{"attributes": a} for a in attrs]}

    (states / "state-05-network.json").write_text(json.dumps({"resources": [
        res("huaweicloud_vpc", {"id": "v-hub", "name": "lz-hub-vpc-dmz-01"},
            {"id": "v-app", "name": "lz-app-vpc-01"}),
        res("huaweicloud_vpc_subnet",
            {"id": "s-nat", "name": "lz-hub-subnet-nat-01", "vpc_id": "v-hub"},
            {"id": "s-elb", "name": "lz-hub-subnet-elb-01", "vpc_id": "v-hub"},
            {"id": "s-cmp", "name": "lz-app-subnet-compute-01", "vpc_id": "v-app"},
            {"id": "s-att", "name": "lz-app-subnet-att-01", "vpc_id": "v-app"}),
        res("huaweicloud_natv3_gateway", {"name": "lz-nat-01", "vpc_id": "v-hub", "subnet_id": "s-nat"}),
        res("huaweicloud_er_vpc_attachment", {"name": "lz-app-attach", "vpc_id": "v-app", "subnet_id": "s-att"}),
        res("huaweicloud_vpc_flow_log", {"resource_id": "v-app"}),
        res("huaweicloud_vpc_route", {"vpc_id": "v-app"}),
    ]}), encoding="utf-8")
    (states / "state-12-workloads.json").write_text(json.dumps({"resources": [
        res("huaweicloud_compute_instance", {"name": "lz-app-ecs-01", "network": [{"uuid": "s-cmp"}],
                                             "tags": {"project": "Payroll", "env": "prd"}}),
    ]}), encoding="utf-8")
    return envs, states


def _run(envs: Path, out: Path, *extra) -> openpyxl.Workbook:
    r = subprocess.run([sys.executable, "-m", "lz_pipeline.tools.gen_ipam", "--envs-dir", str(envs),
                        "--out", str(out), *extra], cwd=PIPELINE, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    return openpyxl.load_workbook(out)


def _remarks(wb):
    notes = {r[4]: r[5] for r in wb["Blocks"].iter_rows(min_row=2, values_only=True) if r[3] == "Allocated"}
    purposes = {r[1]: r[3] for r in wb["Subnets"].iter_rows(min_row=2, values_only=True)}
    return notes, purposes


def test_remarks_describe_what_is_deployed(tmp_path):
    envs, states = _estate(tmp_path)
    notes, purposes = _remarks(_run(envs, tmp_path / "ipam.xlsx", "--states-dir", str(states)))

    assert notes["lz-hub-vpc-dmz-01"] == ("Hub VPC (lz-infra) - internet edge (DMZ). Hosts 1 NAT gateway. "
                                          "1 of 2 subnets still empty (1 for load balancers).")
    assert notes["lz-app-vpc-01"] == ("Spoke VPC (lz-app). Hosts 1 ECS VM (Payroll, prd), 1 ER attachment. "
                                      "Flow log enabled.")
    assert purposes == {
        "lz-hub-subnet-nat-01": "Hosts NAT gateway lz-nat-01.",
        "lz-hub-subnet-elb-01": "Intended for load balancers; nothing deployed yet.",
        "lz-app-subnet-compute-01": "For the application tier. Hosts ECS VM lz-app-ecs-01.",
        "lz-app-subnet-att-01": "Hosts ER attachment lz-app-attach.",
    }


def test_without_states_keeps_role_labels(tmp_path):
    envs, _ = _estate(tmp_path)
    notes, purposes = _remarks(_run(envs, tmp_path / "ipam.xlsx"))
    assert notes == {"lz-hub-vpc-dmz-01": "Hub VPC (lz-infra) - internet edge (DMZ)",
                     "lz-app-vpc-01": "Spoke VPC (lz-app)"}
    assert purposes["lz-app-subnet-att-01"] == "ER attachment"
    assert purposes["lz-app-subnet-compute-01"] is None
