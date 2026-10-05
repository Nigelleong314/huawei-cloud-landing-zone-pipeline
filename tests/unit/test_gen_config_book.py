"""gen_config_book: the spec-driven, state-driven and long-list parts of the book.

A near-empty envs tree plus one state pull and a spec checks that --ir lists
the enabled app-scoped permission sets (they live only in generated HCL, so
tfvars alone cannot show them), that the observability sheet names the CTS
key-event notifications from state, and that long CFW address/domain groups
are no longer cut at 900 characters.
"""

import json
import subprocess
import sys
from pathlib import Path

import openpyxl

PIPELINE = Path(__file__).resolve().parents[2] / "pipeline"


def _cells(ws):
    return [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]


def test_spec_sections_cts_notifications_and_full_groups(tmp_path):
    envs, states = tmp_path / "envs", tmp_path / "states"
    (envs / "09-network-cfw").mkdir(parents=True)
    states.mkdir()
    members = [f"10.0.{i}.0/24" for i in range(120)]
    (envs / "09-network-cfw" / "terraform.tfvars.json").write_text(json.dumps({
        "address_groups": [{"name": "lz-ag-onprem", "members": members}],
        "domain_groups": [{"name": "lz-dg-updates", "type": "url",
                           "domains": [f"mirror{i}.example.com" for i in range(60)]}],
    }), encoding="utf-8")
    (states / "state-06-observability.json").write_text(json.dumps({"resources": [
        {"mode": "managed", "type": "huaweicloud_cts_notification", "name": "key_event",
         "instances": [{"index_key": "iam-changes", "attributes": {}},
                       {"index_key": "delete-events", "attributes": {}}]},
    ]}), encoding="utf-8")
    spec = tmp_path / "spec.json"
    spec.write_text(json.dumps({"sheets": {"03_Identity": {"AppPermissionSets": [
        {"Enabled": "TRUE", "Name": "App-Admin", "Account": "lz-app",
         "EnterpriseProjects": "app-prd-ep", "Description": "App admins"},
        {"Enabled": "FALSE", "Name": "Old-Set", "Account": "lz-app"},
    ]}}}), encoding="utf-8")

    out = tmp_path / "book.xlsx"
    r = subprocess.run([sys.executable, "-m", "lz_pipeline.tools.gen_config_book",
                        "--envs-dir", str(envs), "--states-dir", str(states),
                        "--ir", str(spec), "--out", str(out)],
                       cwd=PIPELINE, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    wb = openpyxl.load_workbook(out)

    identity = _cells(wb["03 Identity"])
    assert "App-scoped permission sets (enterprise-project scoped)" in identity
    assert "App-Admin" in identity and "app-prd-ep" in identity
    assert "Old-Set" not in identity

    obs = list(wb["06 Observability"].iter_rows(values_only=True))
    assert ("CTS key-event notifications", "delete-events, iam-changes") in [row[:2] for row in obs]

    fw = _cells(wb["08 Firewall"])
    assert ", ".join(members) in fw
    assert any(isinstance(c, str) and c.endswith("mirror59.example.com") for c in fw)
